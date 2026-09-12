"""Dump the template's real geometry, colours and fonts.

Designing a slide against remembered numbers is how the AI slide ended up a
carbon copy of the Introduction page. This prints what is actually in the
file so a new layout can be drawn against measured coordinates.

    python docs/portfolio/inspect_template.py [slide ...]

Slide numbers are 1-based; with none given it dumps the whole template.
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

BASE = Path(__file__).resolve().parent
TEMPLATE = BASE / "포트폴리오.pptx"
if not TEMPLATE.exists():
    TEMPLATE = Path("/Users/jung/Desktop/Develop/Data3/docs/portfolio/포트폴리오.pptx")


def colour(fmt) -> str:
    try:
        if fmt.type is None:
            return "-"
        if str(fmt.type).startswith("MSO_THEME_COLOR"):
            return f"theme:{fmt.theme_color}"
        return f"#{fmt.rgb}"
    except Exception:
        return "?"


def fill_of(shape) -> str:
    try:
        fill = shape.fill
        if fill.type is None:
            return "inherit"
        if str(fill.type) == "MSO_FILL_TYPE.BACKGROUND (5)":
            return "none"
        return colour(fill.fore_color)
    except Exception:
        return "?"


def dump(slide, number: int) -> None:
    print(f"\n{'=' * 78}\nSLIDE {number}\n{'=' * 78}")
    for shape in slide.shapes:
        geo = (
            f"L{Emu(shape.left).inches:6.2f} T{Emu(shape.top).inches:6.2f} "
            f"W{Emu(shape.width).inches:6.2f} H{Emu(shape.height).inches:6.2f}"
        )
        kind = str(shape.shape_type).split(" (")[0]
        print(f"\n  {shape.name:18} {kind:22} {geo}  fill={fill_of(shape)}")
        if not shape.has_text_frame:
            continue
        for p_i, para in enumerate(shape.text_frame.paragraphs):
            for run in para.runs:
                font = run.font
                size = f"{font.size.pt:g}pt" if font.size else "inherit"
                print(
                    f"      p{p_i} {size:>8} {font.name or 'inherit':<18}"
                    f" bold={str(font.bold):<5} {colour(font.color):<10}"
                    f" spc={run._r.get_or_add_rPr().get('spc', '0'):>5}"
                    f"  {run.text[:46]!r}"
                )


def main() -> None:
    if not TEMPLATE.exists():
        sys.exit(f"template not found: {TEMPLATE}")
    prs = Presentation(TEMPLATE)
    print(f"{TEMPLATE.name}: {len(prs.slides)} slides, "
          f"{Emu(prs.slide_width).inches:.2f} x {Emu(prs.slide_height).inches:.2f} in")

    wanted = [int(a) for a in sys.argv[1:]] or range(1, len(prs.slides) + 1)
    for number in wanted:
        dump(prs.slides[number - 1], number)


if __name__ == "__main__":
    main()
